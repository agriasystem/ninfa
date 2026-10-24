"""Reading every sheet of a CSV / Excel file into bounded grids.

The importers read ONE table, streamed. To understand a document NINFA first needs to SEE it: every
sheet, the lines above the table, the merged cells, what comes after the last row. So this reader
keeps, for each sheet, the first `SCAN_ROWS` rows and the last `TAIL_ROWS` non-blank rows - never
the whole of a big sheet - and `iter_rows` streams a sheet again when the real import needs every
row.

Safety is the readers' own: the same size, row and column limits, the same zip-bomb guard, formulas
read as their cached value and never calculated, no macros, no external links. A workbook is opened
the cheap, streaming way unless it is small, where merged cells (titles, section rows, two-level
headers) are worth a complete read.
"""

import csv
import io
import warnings
import zipfile
from collections import Counter, deque
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.modules.bookings.errors import BookingErrorCode, BookingImportError
from app.modules.bookings.parsers import MAX_COLUMNS, MAX_ROWS, Cell, detect_file_type
from app.modules.bookings.parsers.csv_parser import _DELIMITERS, _decode

SCAN_ROWS = 400
TAIL_ROWS = 60
MAX_SHEETS = 40
# Merged cells need the complete (non-streaming) read: only for workbooks this small once unpacked.
_COMPLETE_READ_MAX_BYTES = 8 * 1024 * 1024
_MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
# A sheet may carry some title lines above `MAX_ROWS` data rows.
_MAX_SHEET_ROWS = MAX_ROWS + 500


@dataclass(frozen=True, slots=True)
class MergedRange:
    """A merged block of cells, 1-based and inclusive (a title across the table, a group header)."""

    min_row: int
    min_col: int
    max_row: int
    max_col: int


@dataclass(slots=True)
class SheetScan:
    """What the scan keeps of one sheet. Row numbers are the spreadsheet's own (1-based)."""

    name: str
    visible: bool
    # Last row that has anything in it (0: the sheet is empty).
    last_row: int
    max_col: int
    # Rows 1..len(head), blank ones included as empty lists (trailing blank rows are dropped).
    head: list[list[Cell]] = field(default_factory=list)
    # The last non-blank rows when the sheet is longer than the head: (row number, cells).
    tail: list[tuple[int, list[Cell]]] = field(default_factory=list)
    merged: tuple[MergedRange, ...] = ()

    def row(self, number: int) -> list[Cell] | None:
        """The cells of one kept row (blank rows are `[]`), or None when it was not kept."""
        if 1 <= number <= len(self.head):
            return self.head[number - 1]
        for tail_number, cells in self.tail:
            if tail_number == number:
                return cells
        return None

    def covers(self, number: int) -> bool:
        return self.row(number) is not None


@dataclass(slots=True)
class WorkbookScan:
    file_type: str  # "csv" | "xlsx"
    sheets: list[SheetScan]
    delimiter: str | None = None


# --- cells ----------------------------------------------------------------------------------------


def _clean(value: object) -> Cell:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value  # type: ignore[return-value]  # int/float/bool/date/datetime/time/None


def _trim(cells: list[Cell]) -> list[Cell]:
    while cells and cells[-1] is None:
        cells.pop()
    return cells


def _too_big(reason: str, **details: object) -> BookingImportError:
    return BookingImportError(BookingErrorCode.FILE_LIMIT_EXCEEDED, reason, details=dict(details))


# --- one sheet, however it was read ---------------------------------------------------------------


def _scan_rows(
    rows: Iterator[list[Cell]], *, name: str, visible: bool, merged: tuple[MergedRange, ...]
) -> SheetScan:
    head: list[list[Cell]] = []
    tail: deque[tuple[int, list[Cell]]] = deque(maxlen=TAIL_ROWS)
    last_row = 0
    max_col = 0
    for number, raw in enumerate(rows, start=1):
        if number > _MAX_SHEET_ROWS:
            raise _too_big("The sheet has more rows than the supported maximum", max_rows=MAX_ROWS)
        cells = _trim([_clean(value) for value in raw])
        if cells:
            last_row = number
            max_col = max(max_col, len(cells))
        if number <= SCAN_ROWS:
            head.append(cells)
        elif cells:
            tail.append((number, cells))
    del head[last_row:]  # trailing blank rows are not part of the sheet
    return SheetScan(
        name=name,
        visible=visible,
        last_row=last_row,
        max_col=max_col,
        head=head,
        tail=[item for item in tail if item[0] > len(head)],
        merged=merged,
    )


# --- workbooks ------------------------------------------------------------------------------------


def _unpacked_sheet_bytes(content: bytes) -> int:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if sum(info.file_size for info in infos) > _MAX_UNCOMPRESSED_BYTES:
                raise _too_big("The workbook is larger than the supported maximum once unpacked")
            return sum(
                info.file_size for info in infos if info.filename.startswith("xl/worksheets/")
            )
    except zipfile.BadZipFile as exc:
        raise BookingImportError(
            BookingErrorCode.UNREADABLE_FILE, "The file is not a readable .xlsx workbook"
        ) from exc


def _load_xlsx(content: bytes, *, read_only: bool) -> Any:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)  # harmless style quirks of real exports
            return load_workbook(io.BytesIO(content), read_only=read_only, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise BookingImportError(
            BookingErrorCode.UNREADABLE_FILE, "The file is not a readable .xlsx workbook"
        ) from exc


def _merged_of(sheet: Any) -> tuple[MergedRange, ...]:
    ranges = getattr(sheet, "merged_cells", None)
    if ranges is None:  # the streaming reader does not expose them
        return ()
    return tuple(
        MergedRange(r.min_row, r.min_col, r.max_row, r.max_col) for r in list(ranges.ranges)[:500]
    )


def _scan_xlsx(content: bytes) -> list[SheetScan]:
    complete = _unpacked_sheet_bytes(content) <= _COMPLETE_READ_MAX_BYTES
    workbook = _load_xlsx(content, read_only=not complete)
    try:
        scans: list[SheetScan] = []
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            visible = sheet.sheet_state == "visible"
            scans.append(
                _scan_rows(
                    (list(row) for row in sheet.iter_rows(values_only=True)),
                    name=str(sheet.title),
                    visible=visible,
                    merged=_merged_of(sheet) if complete else (),
                )
            )
        return scans
    finally:
        workbook.close()


# --- csv ------------------------------------------------------------------------------------------


def detect_delimiter(lines: list[str]) -> str:
    """The delimiter that splits most of the first lines into the SAME number of fields. Looking at
    many lines (not just the first) is what lets a file with a title above its table be read."""
    sample = [line for line in lines if line.strip()][:40]
    best, best_score = ",", 0
    for delimiter in _DELIMITERS:
        counts: list[int] = []
        for line in sample:
            try:
                counts.append(len(next(csv.reader([line], delimiter=delimiter), [])))
            except csv.Error:
                counts.append(1)
        wide = [count for count in counts if count >= 2]
        if not wide:
            continue
        width, frequency = Counter(wide).most_common(1)[0]
        score = frequency * 1000 + width
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _csv_text(content: bytes, encoding: str | None) -> str:
    text = _decode(content, encoding)
    if "\x00" in text:  # UTF-16/binary decoded as single-byte text: never guess, refuse
        raise BookingImportError(
            BookingErrorCode.UNREADABLE_FILE,
            "The file is not plain UTF-8 or Windows-1252 text (is it UTF-16 or binary?)",
        )
    lines = text.splitlines()
    if lines and lines[0].lower().startswith("sep=") and len(lines[0]) == 5:
        text = text[text.index(lines[0]) + len(lines[0]) :].lstrip("\r\n")
    return text


def _csv_delimiter(text: str, forced: str | None, hint: str | None) -> str:
    return forced or (hint if hint in _DELIMITERS else None) or detect_delimiter(text.splitlines())


def _csv_records(text: str, delimiter: str) -> Iterator[list[str]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        yield from reader
    except csv.Error as exc:
        raise BookingImportError(
            BookingErrorCode.UNREADABLE_FILE, "The file could not be parsed as CSV"
        ) from exc


# --- the public reader ----------------------------------------------------------------------------


def scan_workbook(
    content: bytes,
    filename: str,
    *,
    delimiter: str | None = None,
    encoding: str | None = None,
) -> WorkbookScan:
    """Every sheet of the file as a bounded grid. Raises the readers' own `BookingImportError`."""
    file_type = detect_file_type(filename, content)
    if file_type == "xlsx":
        return WorkbookScan("xlsx", _scan_xlsx(content))
    text = _csv_text(content, encoding)
    used = _csv_delimiter(text, delimiter, None)
    scan = _scan_rows(
        (list[Cell](record) for record in _csv_records(text, used)),
        name="",
        visible=True,
        merged=(),
    )
    return WorkbookScan("csv", [scan], delimiter=used)


def iter_rows(
    content: bytes,
    filename: str,
    sheet: str | None,
    *,
    delimiter: str | None = None,
    encoding: str | None = None,
) -> Iterator[tuple[int, list[Cell]]]:
    """Stream the rows of ONE sheet as (row number, cells), blank rows included, trailing blank
    cells trimmed. Row numbers are the spreadsheet's own, the same ones `scan_workbook` reports."""
    file_type = detect_file_type(filename, content)
    if file_type == "xlsx":
        _unpacked_sheet_bytes(content)
        workbook = _load_xlsx(content, read_only=True)
        try:
            if sheet is None or sheet not in workbook.sheetnames:
                raise BookingImportError(
                    BookingErrorCode.SOURCE_SCHEMA_CHANGED,
                    "The sheet is not in the workbook",
                    details={"reason": "sheet_not_found", "sheet_names": list(workbook.sheetnames)},
                )
            for number, row in enumerate(workbook[sheet].iter_rows(values_only=True), start=1):
                if number > _MAX_SHEET_ROWS:
                    raise _too_big("The sheet has more rows than the supported maximum")
                yield number, _trim([_clean(value) for value in row])
        finally:
            workbook.close()
        return
    text = _csv_text(content, encoding)
    used = _csv_delimiter(text, delimiter, None)
    for number, record in enumerate(_csv_records(text, used), start=1):
        if number > _MAX_SHEET_ROWS:
            raise _too_big("The file has more rows than the supported maximum")
        yield number, _trim([_clean(value) for value in record])


def ensure_columns(width: int) -> None:
    if width > MAX_COLUMNS:
        raise _too_big(
            "The file has more columns than the supported maximum", max_columns=MAX_COLUMNS
        )


__all__ = [
    "MAX_SHEETS",
    "SCAN_ROWS",
    "TAIL_ROWS",
    "MergedRange",
    "SheetScan",
    "WorkbookScan",
    "detect_delimiter",
    "ensure_columns",
    "iter_rows",
    "scan_workbook",
]
