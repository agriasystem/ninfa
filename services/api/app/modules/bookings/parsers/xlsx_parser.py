"""XLSX reader (openpyxl, read-only streaming). No .xls, .xlsm or .ods."""

import io
import warnings
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.modules.bookings.errors import BookingErrorCode, BookingImportError
from app.modules.bookings.mapping import FormatOptions
from app.modules.bookings.parsers import MAX_COLUMNS, Cell, SourceRow, SourceTable, limit_rows

# Guard against decompression bombs: the sheets of a real export are far smaller than this.
_MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024


@contextmanager
def _workbook(content: bytes) -> Iterator[Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > _MAX_UNCOMPRESSED_BYTES:
                raise BookingImportError(
                    BookingErrorCode.FILE_LIMIT_EXCEEDED,
                    "The workbook is larger than the supported maximum once unpacked",
                )
        with warnings.catch_warnings():
            # openpyxl warns about harmless style/extension quirks of real-world exports.
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except BookingImportError:
        raise
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise BookingImportError(
            BookingErrorCode.UNREADABLE_FILE, "The file is not a readable .xlsx workbook"
        ) from exc
    try:
        yield workbook
    finally:
        workbook.close()


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean(value: object) -> Cell:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value  # type: ignore[return-value]  # int/float/bool/date/datetime/time/None


def _has_data(sheet: Any) -> bool:
    return any(not all(_blank(v) for v in row) for row in sheet.iter_rows(values_only=True))


def read_xlsx(content: bytes, options: FormatOptions) -> SourceTable:
    with _workbook(content) as workbook:
        names: list[str] = list(workbook.sheetnames)
        if options.sheet_name is not None:
            if options.sheet_name not in names:
                raise BookingImportError(
                    BookingErrorCode.SOURCE_SCHEMA_CHANGED,
                    "The sheet configured for this data source is not in the workbook",
                    details={"reason": "sheet_not_found", "sheet_names": names},
                )
            chosen = options.sheet_name
        else:
            candidates = [
                sheet.title
                for sheet in workbook.worksheets
                if sheet.sheet_state == "visible" and _has_data(sheet)
            ]
            if not candidates:
                raise BookingImportError(
                    BookingErrorCode.EMPTY_FILE, "The workbook contains no data"
                )
            if len(candidates) > 1:
                raise BookingImportError(
                    BookingErrorCode.MAPPING_REQUIRED,
                    "The workbook has several sheets with data: choose the sheet to import",
                    details={"reason": "sheet_selection_required", "sheet_names": candidates},
                )
            chosen = candidates[0]

        header: list[str] | None = None
        header_number = 0
        for number, row in enumerate(workbook[chosen].iter_rows(values_only=True), start=1):
            if not all(_blank(v) for v in row):
                header = ["" if _blank(v) else str(v).strip() for v in row]
                header_number = number
                break
    if header is None:
        raise BookingImportError(BookingErrorCode.EMPTY_FILE, "The chosen sheet contains no data")
    while header and not header[-1]:
        header.pop()
    if len(header) > MAX_COLUMNS:
        return SourceTable("xlsx", header, chosen, lambda: iter(()))
    width = len(header)

    def rows() -> Iterator[SourceRow]:
        with _workbook(content) as workbook:
            sheet = workbook[chosen]
            for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if number <= header_number:
                    continue
                cells = [_clean(v) for v in row[:width]]
                if all(cell is None for cell in cells):
                    continue
                cells.extend([None] * (width - len(cells)))
                yield SourceRow(number, cells)

    return SourceTable("xlsx", header, chosen, lambda: limit_rows(rows()))
