"""CSV/XLSX structured-file reading, reused from the Gate 2 readers (Gate 8, Part A/E)."""

import io
from datetime import date

import openpyxl
import pytest
from sqlalchemy.orm import Session

from app.modules.labor.errors import LaborErrorCode, LaborImportError
from app.modules.labor.service import LaborImportService
from tests.labor_support import LaborFactory, LaborTenant, labor_tenant

HEADERS = ["work_date", "role", "planned_hours"]
COLUMN_MAPPING = {
    "work_date": {"column": "work_date"},
    "role": {"column": "role"},
    "planned_hours": {"column": "planned_hours"},
}


def _service(session: Session, tenant: LaborTenant) -> LaborImportService:
    return LaborImportService(session, tenant.context)


def _confirm(
    service: LaborImportService, tenant: LaborTenant, headers: list[str] | None = None
) -> None:
    service.save_mapping(
        tenant.labor_data_source.id, headers=headers or HEADERS, column_mapping=COLUMN_MAPPING
    )


def test_semicolon_delimited_csv(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date;role;planned_hours\n2026-03-10;Reception;8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_comma_delimited_csv(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date,role,planned_hours\n2026-03-10,Reception,8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_tab_delimited_csv(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date\trole\tplanned_hours\n2026-03-10\tReception\t8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_utf8_bom_csv(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = "﻿work_date,role,planned_hours\n2026-03-10,Reception,8\n".encode()
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_cp1252_csv(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = "work_date,role,planned_hours\n2026-03-10,Réception,8\n".encode("cp1252")
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def _xlsx_bytes(rows: list[list[object]], sheet_name: str = "Sheet1") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_file(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = _xlsx_bytes([["work_date", "role", "planned_hours"], ["2026-03-10", "Reception", 8]])
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.xlsx",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_duplicate_headers_are_rejected(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date,work_date,role,planned_hours\n2026-03-10,x,Reception,8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert not result.succeeded
    assert result.error_code == LaborErrorCode.DUPLICATE_HEADER


def test_ambiguous_multi_sheet_workbook_requires_mapping(
    db_session: Session, factory: LaborFactory
) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)

    workbook = openpyxl.Workbook()
    workbook.active.title = "March"
    workbook.active.append(["work_date", "role", "planned_hours"])
    workbook.active.append(["2026-03-10", "Reception", 8])
    second = workbook.create_sheet("April")
    second.append(["work_date", "role", "planned_hours"])
    second.append(["2026-04-10", "Reception", 8])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()

    description = service.describe_file(
        tenant.labor_data_source.id, filename="w.xlsx", content=content
    )
    assert description.requires_sheet_selection is True
    assert set(description.sheet_names) == {"March", "April"}

    # Importing without a sheet chosen in the mapping fails the same way.
    service.save_mapping(
        tenant.labor_data_source.id,
        headers=HEADERS,
        column_mapping=COLUMN_MAPPING,
    )
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.xlsx",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert not result.succeeded
    assert result.error_code == LaborErrorCode.MAPPING_REQUIRED


def test_extra_unmapped_columns_are_ignored(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date,role,planned_hours,shift_notes\n2026-03-10,Reception,8,anything at all\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_a_schema_change_that_drops_a_mapped_column_is_detected(
    db_session: Session, factory: LaborFactory
) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date,planned_hours\n2026-03-10,8\n"  # "role" column removed
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert not result.succeeded
    assert result.error_code == LaborErrorCode.SOURCE_SCHEMA_CHANGED


def test_constant_category_is_supported(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    service.save_mapping(
        tenant.labor_data_source.id,
        headers=["work_date", "planned_hours"],
        column_mapping={
            "work_date": {"column": "work_date"},
            "planned_hours": {"column": "planned_hours"},
            "labor_category": {"constant": "MAINTENANCE"},
        },
    )
    content = b"work_date,planned_hours\n2026-03-10,8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_constant_currency_is_supported(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    service.save_mapping(
        tenant.labor_data_source.id,
        headers=["work_date", "role", "planned_hours", "planned_cost"],
        column_mapping={
            "work_date": {"column": "work_date"},
            "role": {"column": "role"},
            "planned_hours": {"column": "planned_hours"},
            "planned_cost": {"column": "planned_cost"},
            "currency": {"constant": "EUR"},
        },
    )
    content = b"work_date,role,planned_hours,planned_cost\n2026-03-10,Reception,8,100\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_constant_work_date_is_prohibited(db_session: Session, factory: LaborFactory) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    with pytest.raises(LaborImportError) as info:
        service.save_mapping(
            tenant.labor_data_source.id,
            headers=["role", "planned_hours"],
            column_mapping={
                "work_date": {"constant": "2026-03-10"},
                "role": {"column": "role"},
                "planned_hours": {"column": "planned_hours"},
            },
        )
    assert info.value.error_code == LaborErrorCode.INVALID_MAPPING


def test_ambiguous_date_format_is_rejected_without_a_configured_pattern(
    db_session: Session, factory: LaborFactory
) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    content = b"work_date,role,planned_hours\n01/02/2026,Reception,8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert not result.succeeded
    assert result.error_code == LaborErrorCode.AMBIGUOUS_DATE_FORMAT


def test_a_configured_date_format_resolves_the_ambiguity(
    db_session: Session, factory: LaborFactory
) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    service.save_mapping(
        tenant.labor_data_source.id,
        headers=HEADERS,
        column_mapping=COLUMN_MAPPING,
        format_options={"date_formats": {"work_date": "%d/%m/%Y"}},
    )
    content = b"work_date,role,planned_hours\n01/02/2026,Reception,8\n"
    result = service.import_file(
        tenant.labor_data_source.id,
        filename="w.csv",
        content=content,
        snapshot_local_date=date(2026, 3, 1),
    )
    assert result.succeeded


def test_a_confirmed_mapping_profile_is_reused_across_imports(
    db_session: Session, factory: LaborFactory
) -> None:
    tenant = labor_tenant(factory)
    service = _service(db_session, tenant)
    _confirm(service, tenant)
    for day in ("2026-03-10", "2026-03-11"):
        content = f"work_date,role,planned_hours\n{day},Reception,8\n".encode()
        result = service.import_file(
            tenant.labor_data_source.id,
            filename="w.csv",
            content=content,
            snapshot_local_date=date(2026, 3, int(day[-2:])),
        )
        assert result.succeeded
