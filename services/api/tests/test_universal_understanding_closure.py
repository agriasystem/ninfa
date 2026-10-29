"""Universal Understanding 2.1 - closing two targets, each rule with its counter-example.

* the lines above a table are read by what their LABEL says they are (a department, a note, a
  customer, a person...), never by the value alone;
* the document date and the registration date of a register are not two candidates for one field;
* the words of a status are asked as words, not twice;
* the scorer counts a column the person has just confirmed as right (a gap of `_extended`).

No database, no network, no model.
"""

from datetime import date, timedelta
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook

from app.modules.document_understanding.context import LabelKind, kind_of_label, read_line
from app.modules.document_understanding.titles import department_words, stated_department_words
from tests.benchmarks.smart_import_v3.corpus import build_corpus
from tests.benchmarks.smart_import_v3.model import Golden, Source
from tests.benchmarks.smart_import_v3.observe import Observed, OClaim, OQuestion, make_service, run
from tests.benchmarks.smart_import_v3.score import score_case

# --- the lines above a table ---------------------------------------------------------------------


@pytest.mark.parametrize(
    ("line", "kind"),
    [
        ("Reparto: Spa", LabelKind.DEPARTMENT),
        ("Department: Housekeeping", LabelKind.DEPARTMENT),
        ("Settore: Sala", LabelKind.DEPARTMENT),
        ("Periodo: Ottobre 2026", LabelKind.PERIOD),
        ("Struttura: Hotel Centro", LabelKind.PROPERTY),
        ("Responsabile: Mario Rossi", LabelKind.PERSON),
        ("Cliente: Spa S.r.l.", LabelKind.PARTY),
        ("Nota: Spa chiusa lunedi", LabelKind.NOTE),
        ("Turni: Housekeeping", None),  # a label nobody knows
        ("TURNI HOUSEKEEPING", None),  # not a label/value pair at all: a title
    ],
)
def test_a_line_above_a_table_is_known_by_its_label(line: str, kind: LabelKind | None) -> None:
    assert read_line(line).kind == kind


def test_only_the_label_makes_a_word_the_department() -> None:
    assert department_words(["Reparto: Spa"]) == ["Spa"]
    assert department_words(["Department: Housekeeping", "Periodo: Ottobre 2026"]) == [
        "Housekeeping"
    ]
    # the same word after the label of a note, a customer, a person, a hotel: not a department
    for line in (
        "Nota: Spa chiusa lunedi",
        "Cliente: Spa S.r.l.",
        "Responsabile: Direzione",
        "Struttura: Hotel Spa Centro",
    ):
        assert department_words([line]) == [], line
        assert stated_department_words([line]) == [], line


def test_a_stated_department_beats_every_other_line_and_a_title_is_still_a_title() -> None:
    block = ["Reparto: Sala", "Nota: Spa chiusa", "Responsabile: Direzione"]

    assert department_words(block) == ["Sala"]
    assert department_words(["TURNI SPA"]) == ["SPA"]  # a title: as before
    assert department_words(["Turni: Housekeeping"]) == ["Housekeeping"]  # unknown label: as before
    assert stated_department_words(["TURNI SPA"]) == []  # (a title is read, not stated)


def test_the_label_decides_in_any_order_of_words() -> None:
    assert kind_of_label("Nota cliente") == LabelKind.NOTE
    assert kind_of_label("Reparto di appartenenza") == LabelKind.DEPARTMENT
    assert kind_of_label("Boh") is None


def _roster(block: list[tuple[str, str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for label, value in block:
        sheet.append([label, value])
    sheet.append([])
    sheet.append(["Data", "Operatore", "Ore previste", "Ore effettive"])
    start = date(2026, 10, 19)
    for i in range(8):
        sheet.append([start + timedelta(days=i % 4), f"Persona {i % 3}", 7 + i % 2, 7 + i % 2])
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _observe(block: list[tuple[str, str]]) -> Observed:
    return run(make_service(), "labor", "turni.xlsx", _roster(block))


def test_a_department_stated_above_the_table_is_the_department_of_its_rows() -> None:
    seen = _observe([("Reparto:", "SPA E BENESSERE"), ("Periodo:", "19-25 ottobre 2026")])

    assert seen.level == "READY" and not seen.questions
    assert any(c.field == "role" and c.kind == "TITLE" for c in seen.claims)


@pytest.mark.parametrize(
    "block",
    [
        [("Nota:", "Spa chiusa lunedi")],
        [("Cliente:", "Spa S.r.l.")],
        [("Responsabile:", "Direzione")],
        [("Struttura:", "Hotel Centro"), ("Periodo:", "Ottobre 2026")],
    ],
)
def test_the_same_words_under_another_label_are_not_a_department(
    block: list[tuple[str, str]],
) -> None:
    seen = _observe(block)

    assert not any(c.field == "role" for c in seen.claims)
    assert any(q.kind == "DEPARTMENT_ROWS" for q in seen.questions)


def test_a_block_of_labels_does_not_hide_the_department_of_a_note_beside_it() -> None:
    seen = _observe([("Reparto:", "Sala"), ("Nota:", "Spa chiusa lunedi")])

    assert any(c.field == "role" and c.kind == "TITLE" for c in seen.claims)


# --- a register with a document date and a registration date --------------------------------------


def _register(headers: list[str]) -> Observed:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Fornitore", "Fattura n.", *headers, "Descrizione", "Imponibile"])
    start = date(2026, 10, 1)
    for i in range(8):
        row: list[Any] = [f"Alfa {i % 3} Srl", f"FT/{100 + i}"]
        row += [start + timedelta(days=i + 2 * k) for k in range(len(headers))]
        row += [f"Fornitura {i}", 100.0 + i]
        sheet.append(row)
    out = BytesIO()
    book.save(out)
    return run(make_service(), "costs", "registro.xlsx", out.getvalue())


def test_the_document_date_is_the_invoice_date_and_the_registration_date_is_no_contest() -> None:
    seen = _register(["Data doc.", "Data reg."])

    dates = [c for c in seen.claims if c.field == "invoice_date"]
    assert [(c.header, c.level) for c in dates] == [("Data doc.", "HIGH")]
    assert not any(q.kind == "COLUMN_MEANING" for q in seen.questions)


def test_a_register_with_only_a_registration_date_still_has_an_invoice_date() -> None:
    seen = _register(["Data registrazione"])

    assert [(c.header, c.level) for c in seen.claims if c.field == "invoice_date"] == [
        ("Data registrazione", "HIGH")
    ]


def test_two_dates_that_both_say_invoice_stay_a_choice() -> None:
    seen = _register(["Data fattura", "Data emissione"])

    assert not any(c.field == "invoice_date" and c.level == "HIGH" for c in seen.claims)


# --- the words of a status are asked as words, not twice -----------------------------------------


def test_a_status_written_as_codes_is_asked_once() -> None:
    document = next(d for d in build_corpus() if d.id.startswith("BO09"))

    seen = run(make_service(), document.card, document.filename, document.content)

    kinds = [(q.kind, q.column) for q in seen.questions]
    assert ("VALUE_LABELS", None) in kinds
    assert not any(kind == "COLUMN_MEANING" and column == "Stato" for kind, column in kinds)


# --- the scorer: a column the person has just confirmed is right --------------------------------


def _golden(amount_header: str, claimed_column: str) -> tuple[Golden, Observed, Observed]:
    golden = Golden(
        "costs",
        "ASK",
        fields={},
        ambiguous=("net_amount",),
        questions=("AMOUNT_BASIS",),
        answers={
            "amount_basis": "NET",
            "amount_column": {"header": [amount_header], "occurrence": 1},
        },
    )
    asked = OQuestion("AMOUNT_BASIS", amount_header, "line_total", None, None, (), ())
    first = Observed(
        level="ASK",
        questions=[asked],
        claims=[OClaim("line_total", "COLUMN", claimed_column, "MEDIUM")],
    )
    final = Observed(level="READY", claims=[OClaim("line_total", "COLUMN", claimed_column, "HIGH")])
    return golden, first, final


def _entry() -> dict[str, Any]:
    return {"id": "X", "card": "costs", "filename": "x.xlsx", "tags": [], "sheets": 1}


def test_the_amount_the_person_confirms_as_net_is_a_right_claim() -> None:
    golden, first, final = _golden("Importo", "Importo")

    result = score_case(_entry(), golden, first, [final])

    assert result.required.claims == 1 and result.required.claims_correct == 1
    assert result.required.golden == 1 and result.required.correct == 1


def test_but_another_column_claimed_as_the_net_is_still_wrong() -> None:
    golden, first, final = _golden("Importo", "Totale")

    result = score_case(_entry(), golden, first, [final])

    assert result.required.claims == 1 and result.required.claims_correct == 0


def test_without_the_persons_answer_a_claim_on_an_ambiguous_field_is_still_not_right() -> None:
    golden, first, _final = _golden("Importo", "Importo")
    golden.answers = {}

    result = score_case(_entry(), golden, first, None)

    assert result.required.claims == 1 and result.required.claims_correct == 0


def test_the_source_of_a_column_is_a_column() -> None:
    assert Source("COLUMN", ("Importo",), 1).kind == "COLUMN"


def test_once_the_person_has_taught_the_status_words_the_status_column_is_not_asked() -> None:
    from tests.benchmarks.smart_import_v3.dialog import converse

    document = next(d for d in build_corpus() if d.id.startswith("BO09"))
    service = make_service()
    first = run(service, document.card, document.filename, document.content)

    rounds = converse(service, document, first)

    assert [q.kind for q in first.questions] == ["VALUE_LABELS"]
    assert [r.level for r in rounds] == ["READY"]  # one answer, and nothing else is asked
    assert not rounds[0].questions


@pytest.mark.parametrize(
    "headers",
    [
        ["Data documento", "Data registrazione"],
        ["Data registrazione", "Data documento"],  # (the registration date comes first)
        ["Data reg.", "Data doc."],
        ["Registration date", "Document date"],
    ],
)
def test_the_document_date_wins_whichever_column_comes_first(headers: list[str]) -> None:
    seen = _register(headers)

    dates = [c for c in seen.claims if c.field == "invoice_date"]
    assert len(dates) == 1 and dates[0].level == "HIGH"
    assert dates[0].header in {"Data documento", "Data doc.", "Document date"}


def test_a_net_that_says_taxable_is_the_net() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Supplier", "Invoice No.", "Invoice date", "Description", "Taxable Amount"])
    for i in range(8):
        sheet.append(
            [f"Alfa {i % 3} Ltd", f"FT/{i}", date(2026, 10, 1 + i), f"Item {i}", 100.0 + i]
        )
    out = BytesIO()
    book.save(out)

    seen = run(make_service(), "costs", "ledger.xlsx", out.getvalue())

    assert any(c.field == "line_total" and c.level == "HIGH" for c in seen.claims)
    assert not any(q.kind == "AMOUNT_BASIS" for q in seen.questions)
