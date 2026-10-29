"""The corpus as a whole: assembled, checked against its own files, summarised, hashed.

`validate` is the benchmark's check on ITSELF. It reopens every generated file with openpyxl / the
csv module - never with Smart Import's readers - and verifies what the truth claims: the header
cells are where the truth says, the first record is where the truth says, the personal values are
really in the file. A corpus that fails this is a bug of the benchmark, not a finding.
"""

import csv
import hashlib
import io
from collections import Counter
from typing import Any

from openpyxl import load_workbook

from tests.benchmarks.smart_import_v3.corpus_bookings import BOOKINGS
from tests.benchmarks.smart_import_v3.corpus_costs import COSTS
from tests.benchmarks.smart_import_v3.corpus_labor import LABOR
from tests.benchmarks.smart_import_v3.corpus_office import (
    BOOKINGS_OFFICE,
    COSTS_OFFICE,
    LABOR_OFFICE,
)
from tests.benchmarks.smart_import_v3.formats_truth import BY_CARD, FORMATS
from tests.benchmarks.smart_import_v3.match import norm
from tests.benchmarks.smart_import_v3.model import Doc, golden_to_json
from tests.benchmarks.smart_import_v3.oracle_review import apply_review

DOMAIN_NAMES = {"labor": "Personale", "bookings": "Prenotazioni", "costs": "Costi"}


def build_corpus() -> list[Doc]:
    docs = [
        *LABOR.build(),
        *LABOR_OFFICE.build(),
        *BOOKINGS.build(),
        *BOOKINGS_OFFICE.build(),
        *COSTS.build(),
        *COSTS_OFFICE.build(),
    ]
    ids = Counter(d.id for d in docs)
    duplicated = [i for i, n in ids.items() if n > 1]
    if duplicated:
        raise ValueError(f"duplicate document ids: {duplicated}")
    apply_formats(docs)
    apply_review(docs)
    return docs


def apply_formats(docs: list[Doc]) -> None:
    """Adds how the text dates and amounts of some documents are written (`formats_truth`)."""
    used: set[str] = set()
    for doc in docs:
        slug = doc.id.split("_", 1)[1]
        key = BY_CARD.get((doc.card, slug), slug)
        if key in FORMATS:
            doc.golden.date_order, doc.golden.decimal = FORMATS[key]
            used.add(key)
    unused = set(FORMATS) - used
    if unused:
        raise ValueError(f"formats for documents that do not exist: {sorted(unused)}")


def corpus_hash(docs: list[Doc]) -> str:
    digest = hashlib.sha256()
    for doc in sorted(docs, key=lambda d: d.id):
        digest.update(doc.id.encode())
        digest.update(doc.content)
    return digest.hexdigest()


def manifest(docs: list[Doc]) -> list[dict[str, Any]]:
    """The truth, one entry per document, as plain JSON (the only thing the scorer reads)."""
    return [
        {
            "id": d.id,
            "card": d.card,
            "filename": d.filename,
            "bytes": len(d.content),
            "sha256": hashlib.sha256(d.content).hexdigest(),
            "tags": list(d.tags),
            "structure": list(d.structure),
            "sheets": d.sheet_count,
            "golden": golden_to_json(d.golden),
        }
        for d in docs
    ]


def summary(docs: list[Doc]) -> dict[str, Any]:
    by_card = Counter(d.card for d in docs)
    tags = Counter(t for d in docs for t in d.tags)
    shapes = {tuple(sorted(d.structure)) for d in docs}
    return {
        "documents": len(docs),
        "by_card": {DOMAIN_NAMES[k]: v for k, v in by_card.items()},
        "xlsx": tags["xlsx"],
        "csv": tags["csv"],
        "negative_refusals": sum(
            1 for d in docs if d.golden.outcome in {"REJECT", "NOT_AUTOMATIC", "WARN"}
        ),
        "negative_tagged": tags["negative"],
        "adversarial": tags["adversarial"],
        "office_realistic": tags["office"],
        "large": tags["large"],
        "with_personal_data": sum(1 for d in docs if d.golden.pii_columns),
        "multi_sheet": sum(1 for d in docs if d.sheet_count >= 2),
        "distinct_structures": len(shapes),
        "outcomes": dict(Counter(d.golden.outcome for d in docs)),
        "sha256": corpus_hash(docs),
    }


# --- the corpus checks itself ---------------------------------------------------------------------


def _rows_xlsx(content: bytes, sheet: str | None) -> list[list[str]]:
    book = load_workbook(io.BytesIO(content), data_only=True)
    ws = book[sheet] if sheet else book.worksheets[0]
    return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _rows_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    dialect = ";" if text.count(";") > text.count(",") else ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=dialect)]


def _whole_file_text(doc: Doc) -> str:
    """Every cell of every sheet (including hidden ones) as one string."""
    if doc.filename.endswith(".csv"):
        return " ".join(" ".join(row) for row in _rows_csv(doc.content))
    book = load_workbook(io.BytesIO(doc.content), data_only=True)
    return " ".join(
        " ".join("" if c is None else str(c) for c in row)
        for ws in book.worksheets
        for row in ws.iter_rows(values_only=True)
    )


def _row_text(row: list[str]) -> str:
    return " | ".join(norm(cell) for cell in row)


def validate(docs: list[Doc]) -> list[str]:
    """Problems of the benchmark itself (empty when the truth matches the files)."""
    problems: list[str] = []
    for doc in docs:
        g = doc.golden
        try:
            sheet = g.sheets[0] if g.sheets else None
            rows = (
                _rows_csv(doc.content)
                if doc.filename.endswith(".csv")
                else _rows_xlsx(doc.content, sheet)
            )
        except Exception as error:  # a file the benchmark cannot even reopen
            problems.append(f"{doc.id}: cannot reopen ({type(error).__name__}: {error})")
            continue
        if g.header_rows and all(n <= len(rows) for n in g.header_rows):
            bottom = _row_text(rows[g.header_rows[-1] - 1])
            top = _row_text(rows[g.header_rows[0] - 1])
            wanted = [s.header for s in g.fields.values() if s.kind == "COLUMN"]
            wanted += [*g.unmapped, *g.pii_columns]
            for header in wanted:
                if norm(header[-1]) not in bottom:
                    problems.append(
                        f"{doc.id}: header cell {header!r} not on row {g.header_rows[-1]}"
                    )
                elif len(header) > 1 and len(g.header_rows) > 1 and norm(header[0]) not in top:
                    problems.append(
                        f"{doc.id}: upper header {header[0]!r} not on row {g.header_rows[0]}"
                    )
        elif g.header_rows:
            problems.append(f"{doc.id}: header rows {g.header_rows} are beyond the sheet")
        if (
            g.first_row is not None
            and g.data_rows
            and (g.first_row > len(rows) or not any(rows[g.first_row - 1]))
        ):
            problems.append(f"{doc.id}: first record row {g.first_row} is empty")
        if g.pii_values:
            blob = _whole_file_text(doc).casefold()
            missing = [v for v in g.pii_values if v.casefold() not in blob]
            if missing:
                problems.append(f"{doc.id}: personal values not in the file: {missing[:2]}")
        for outcome in (g.outcome,):
            if outcome not in {"READY", "ASK", "REJECT", "NOT_AUTOMATIC", "WARN"}:
                problems.append(f"{doc.id}: unknown outcome {outcome}")
        if g.outcome == "ASK" and not g.questions:
            problems.append(f"{doc.id}: ASK without a question")
    return problems
