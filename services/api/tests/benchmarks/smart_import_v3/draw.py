"""Draws the documents of the corpus as REAL bytes (an .xlsx written with openpyxl, a CSV the way
Italian or English Excel writes it) and keeps track of where everything was put.

The row numbers of a header, of the first record, the count of records and of the lines that are
not records are recorded WHILE drawing (`TableRef`), so the golden truth knows them without looking
at what any reader later makes of the file.

Output is deterministic: fixed workbook properties and a re-zipped container with fixed member
dates, so the same corpus has the same hash on every run.
"""

import csv
import io
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_FIXED_TIME = datetime(2026, 10, 1, 9, 0, 0)
_ZIP_TIME = (2026, 10, 1, 9, 0, 0)

# --- cells ----------------------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Styled:
    """A value and how to show it. `text=True` writes a string even if it looks like a formula."""

    value: Any
    fmt: str | None = None  # Excel number format
    style: str | None = None  # a named style
    text: bool = False


@dataclass(frozen=True, slots=True)
class Section:
    """A line of its own that labels the records under it ("RECEPTION")."""

    text: str


@dataclass(frozen=True, slots=True)
class Subtotal:
    cells: Sequence[Any]


@dataclass(frozen=True, slots=True)
class Total:
    cells: Sequence[Any]


@dataclass(frozen=True, slots=True)
class Note:
    text: str


@dataclass(frozen=True, slots=True)
class TableRef:
    """Where a table ended up on its sheet."""

    sheet: str
    header_rows: tuple[int, ...]
    first_row: int
    last_row: int  # the last record
    data_rows: int
    skipped: dict[str, int] = field(default_factory=dict)

    def merged(self, *others: "TableRef") -> "TableRef":
        """Several sheets that are one table (a partition): the records add up."""
        every = (self, *others)
        totals = {"totals": 0, "notes": 0, "sections": 0}
        for ref in every:
            for key, count in ref.skipped.items():
                totals[key] += count
        return TableRef(
            self.sheet,
            self.header_rows,
            self.first_row,
            self.last_row,
            sum(ref.data_rows for ref in every),
            totals,
        )


# --- a sheet --------------------------------------------------------------------------------------


class Sheet:
    def __init__(
        self,
        title: str,
        *,
        hidden: bool = False,
        left: int = 0,
        widths: dict[int, float] | None = None,
        freeze: str | None = None,
    ) -> None:
        self.title = title
        self.hidden = hidden
        self.left = left  # empty columns before the first one
        self.widths = widths or {}
        self.freeze = freeze
        self.rows: list[list[Any]] = []
        self.styles: list[str | None] = []
        self.merges: list[tuple[int, int, int, int]] = []

    @property
    def last_row(self) -> int:
        return len(self.rows)

    def add(self, cells: Sequence[Any] = (), *, style: str | None = None) -> int:
        self.rows.append(list(cells))
        self.styles.append(style)
        return len(self.rows)

    def blank(self, count: int = 1) -> int:
        for _ in range(count):
            self.add(())
        return self.last_row

    def merge(self, row: int, first_col: int, last_col: int) -> None:
        """Merge columns first_col..last_col (1-based, before the left offset) of a row."""
        self.merges.append((row, first_col, row, last_col))

    def merge_range(self, first_row: int, first_col: int, last_row: int, last_col: int) -> None:
        """Merge a block of cells (a department written once for several rows)."""
        self.merges.append((first_row, first_col, last_row, last_col))

    def put(self, row: int, col: int, value: Any, *, style: str | None = None) -> None:
        """Write a cell into an existing row (1-based column, before the left offset): a block
        drawn beside the table."""
        cells = self.rows[row - 1]
        while len(cells) < col:
            cells.append(None)
        cells[col - 1] = Styled(value, style=style) if style else value

    def title_line(self, text: str, width: int, *, style: str = "title") -> int:
        row = self.add([text], style=style)
        if width > 1:
            self.merge(row, 1, width)
        return row

    def pairs(self, items: Sequence[tuple[str, Any]], *, style: str | None = "label") -> int:
        """Label / value lines ("Hotel", "Hotel Aurora"). Returns the last row."""
        for label, value in items:
            self.add([label, value], style=style)
        return self.last_row


def write_table(
    sheet: Sheet,
    header: Sequence[str],
    body: Sequence[Any],
    *,
    upper: Sequence[str | None] | None = None,
    upper_merges: Sequence[tuple[int, int]] = (),
    header_style: str | None = "header",
    row_style: str | None = None,
) -> TableRef:
    """Writes a header (optionally with an upper level) and a body, and returns where it is.

    A body item is a record (a list of cells) or one of `Section`, `Subtotal`, `Total`, `Note`;
    only a record counts as data."""
    width = len(header)
    header_rows: list[int] = []
    if upper is not None:
        header_rows.append(sheet.add(upper, style="header_top"))
        for first, last in upper_merges:
            sheet.merge(sheet.last_row, first, last)
    header_rows.append(sheet.add(header, style=header_style))
    first = 0
    last = 0
    data = 0
    skipped = {"totals": 0, "notes": 0, "sections": 0}
    for item in body:
        if isinstance(item, Section):
            row = sheet.add([item.text], style="section")
            sheet.merge(row, 1, width)
            skipped["sections"] += 1
        elif isinstance(item, (Subtotal, Total)):
            sheet.add(item.cells, style="total")
            skipped["totals"] += 1
        elif isinstance(item, Note):
            sheet.add([item.text], style="note")
            skipped["notes"] += 1
        else:
            row = sheet.add(item, style=row_style)
            data += 1
            first = first or row
            last = row
    return TableRef(sheet.title, tuple(header_rows), first, last, data, skipped)


# --- xlsx -----------------------------------------------------------------------------------------

_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _apply_style(cell: Any, name: str | None) -> None:
    if name is None:
        return
    if name == "title":
        cell.font = Font(bold=True, size=14)
    elif name == "subtitle":
        cell.font = Font(italic=True, size=11, color="555555")
    elif name == "header":
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BOX
    elif name == "header_top":
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="BDD7EE")
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BOX
    elif name == "header_soft":
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7E6E6")
        cell.border = _BOX
    elif name == "total":
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
        cell.border = Border(top=Side(style="medium"))
    elif name == "section":
        cell.font = Font(bold=True, color="1F4E78")
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    elif name == "note":
        cell.font = Font(italic=True, color="7F7F7F")
    elif name == "label":
        cell.font = Font(bold=True)
    elif name == "warning":
        cell.font = Font(bold=True, color="C00000")
        cell.fill = PatternFill("solid", fgColor="FFE699")
    elif name == "band":
        cell.fill = PatternFill("solid", fgColor="F2F2F2")


def _write_cell(ws: Any, row: int, col: int, raw: Any, row_style: str | None) -> None:
    value, fmt, style, text = raw, None, row_style, False
    if isinstance(raw, Styled):
        value, fmt, text = raw.value, raw.fmt, raw.text
        style = raw.style or row_style
    if value is None:
        if style is not None:
            _apply_style(ws.cell(row=row, column=col), style)
        return
    cell = ws.cell(row=row, column=col, value=value)
    if text:
        cell.data_type = "s"
    if fmt is not None:
        cell.number_format = fmt
    elif isinstance(value, datetime):
        cell.number_format = (
            "DD/MM/YYYY" if value.time() == datetime.min.time() else "DD/MM/YYYY HH:MM"
        )
    elif isinstance(value, date):
        cell.number_format = "DD/MM/YYYY"
    _apply_style(cell, style)


def _normalize_zip(data: bytes) -> bytes:
    """The same workbook always gives the same bytes (member dates fixed, members in order)."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            data = source.read(name)
            if name == "docProps/core.xml":
                # openpyxl stamps the save time into the file: fix it
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-10-01T09:00:00Z\g<2>",
                    data,
                )
            target.writestr(info, data)
    return out.getvalue()


def xlsx_bytes(sheets: Sequence[Sheet]) -> bytes:
    book = Workbook()
    book.properties.creator = "benchmark"
    book.properties.created = _FIXED_TIME
    book.properties.modified = _FIXED_TIME
    first = True
    for sheet in sheets:
        ws = book.active if first else book.create_sheet()
        first = False
        ws.title = sheet.title
        if sheet.hidden:
            ws.sheet_state = "hidden"
        for index, (cells, style) in enumerate(zip(sheet.rows, sheet.styles, strict=True), start=1):
            for offset, raw in enumerate(cells):
                _write_cell(ws, index, sheet.left + offset + 1, raw, style)
            if not cells and style is not None:
                _apply_style(ws.cell(row=index, column=sheet.left + 1), style)
        for r1, c1, r2, c2 in sheet.merges:
            ws.merge_cells(
                start_row=r1,
                start_column=sheet.left + c1,
                end_row=r2,
                end_column=sheet.left + c2,
            )
        widest = max((len(r) for r in sheet.rows), default=0)
        for col in range(1, widest + 1):
            ws.column_dimensions[get_column_letter(sheet.left + col)].width = sheet.widths.get(
                col, 16
            )
        if sheet.freeze:
            ws.freeze_panes = sheet.freeze
    buffer = io.BytesIO()
    book.save(buffer)
    return _normalize_zip(buffer.getvalue())


# --- csv ------------------------------------------------------------------------------------------


def _csv_cell(value: Any, *, decimal: str, dates: str) -> str:
    if isinstance(value, Styled):
        value = value.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            return (
                value.strftime("%Y-%m-%d %H:%M")
                if dates == "iso"
                else value.strftime("%d/%m/%Y %H:%M")
            )
        return _csv_date(value, dates)
    if isinstance(value, date):
        return _csv_date(value, dates)
    if isinstance(value, float):
        text = repr(value)
        return text.replace(".", decimal) if decimal != "." else text
    return str(value)


def _csv_date(value: date, style: str) -> str:
    if style == "iso":
        return value.strftime("%Y-%m-%d")
    if style == "us":
        return value.strftime("%m/%d/%Y")
    if style == "dotted":
        return value.strftime("%d.%m.%Y")
    return value.strftime("%d/%m/%Y")


def csv_bytes(
    rows: Sequence[Sequence[Any]],
    *,
    delimiter: str = ";",
    decimal: str = ",",
    dates: str = "it",
    encoding: str = "utf-8-sig",
    newline: str = "\r\n",
) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator=newline)
    for row in rows:
        writer.writerow([_csv_cell(cell, decimal=decimal, dates=dates) for cell in row])
    return buffer.getvalue().encode(encoding)


def sheet_csv(sheet: Sheet, **options: Any) -> bytes:
    """A sheet written as a CSV: the left offset becomes empty cells, merges are lost."""
    pad: list[Any] = [None] * sheet.left
    return csv_bytes([[*pad, *row] for row in sheet.rows], **options)
