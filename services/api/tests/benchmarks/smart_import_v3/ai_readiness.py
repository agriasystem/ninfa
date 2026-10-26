"""python -m tests.benchmarks.smart_import_v3.ai_readiness            (from services/api)

AI readiness benchmark: is it safe to switch a model on, and what would it be worth?

The same 240 documents (and a second population of 'unreadable header' variants of them, built so
that the model is actually asked) go through the real Smart Import with a STAND-IN for the model:

* an ORACLE that always answers what the truth says - the most a model could ever be worth;
* models that are WRONG in every way (plainly wrong, wrong and sure of it, wrong in a way that
  passes the value checks, obeying a sentence in a cell, saying it is another kind of data);
* models that FAIL (timeout, 429, 500, garbage, nothing, an answer that weighs a book).

Nothing leaves the process and no vendor is involved. The numbers that decide are: with a wrong or
failing model, does anything change that could cost the hotel (false automation, a file made ready
by the model, a personal value in what the model is shown)? And with a perfect one, how many
questions would really go away?
"""

import argparse
import json
import statistics
import sys
from collections import Counter
from collections.abc import Callable, Mapping
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook

from app.core.tenant import TenantContext
from app.modules.ai.import_understanding.egress import MAX_PAYLOAD_CHARS, EgressGuard
from app.modules.ai.import_understanding.resolver import ImportUnderstandingResolver
from app.modules.ai.import_understanding.types import (
    AiResolution,
    ImportUnderstandingRequest,
    ResolutionReason,
    UnderstandingEvidence,
)
from app.modules.ai.import_understanding.validator import FitOf
from app.modules.data_sources.semantics_bookings import BOOKINGS
from app.modules.data_sources.semantics_costs import COSTS
from app.modules.data_sources.semantics_labor import LABOR
from app.modules.data_sources.smart_import import SmartImportService
from app.modules.data_sources.understanding import UnderstandingService
from app.modules.document_understanding.pii_text import normalize
from tests.benchmarks.smart_import_v3.corpus import build_corpus, manifest
from tests.benchmarks.smart_import_v3.match import header_matches
from tests.benchmarks.smart_import_v3.metrics import ratio
from tests.benchmarks.smart_import_v3.model import (
    CRITICAL_FIELDS,
    Doc,
    Golden,
    Source,
    golden_from_json,
)
from tests.benchmarks.smart_import_v3.observe import Observed, run
from tests.benchmarks.smart_import_v3.score import CaseResult, score_case
from tests.import_ai_support import Scenario, ScenarioProvider, audit

REPO = Path(__file__).resolve().parents[5]
DEFAULT_JSON = REPO / "services/api/tests/benchmarks/results/smart-import-v3-ai-readiness-v1.json"
DEFAULT_MD = REPO / "docs/benchmarks/smart-import-v3-ai-readiness.md"
BASELINE = "1479a8d92e03f3eba8603aaa4a468bdd08227b3e"

SEMANTICS = {"labor": LABOR, "bookings": BOOKINGS, "costs": COSTS}
# the two spellings the importers accept for one business fact (see `model`)
PSEUDO = {
    "department": ("role", "labor_category"),
    "net_amount": ("line_total", "invoice_net_amount"),
}
# the header an obfuscated variant gets: no rule can read it, a person (or a model) can tell from
# the values what it is
UNREADABLE = "Colonna 9"


# --- the stand-ins ---


def _offered(document: Mapping[str, Any], domain: str) -> list[str]:
    return list(document["allowed"]["fields"].get(domain, []))


def truth_fields(golden: Golden, offered: list[str]) -> dict[str, str]:
    """Header cell -> the field a model would be right to name (only fields that were offered)."""
    out: dict[str, str] = {}
    for name, source in golden.fields.items():
        if source.kind != "COLUMN" or not source.header:
            continue
        for candidate in (name, *PSEUDO.get(name, ())):
            if candidate in offered:
                out[source.header[-1]] = candidate
                break
    return out


@dataclass
class Oracle:
    """Answers what the truth says: the best a model could ever be."""

    golden: Golden
    requests: list[ImportUnderstandingRequest] = field(default_factory=list)
    violations: list[str] = field(default_factory=list)
    answered: list[int] = field(default_factory=list)

    def suggest(self, request: ImportUnderstandingRequest) -> object:
        self.requests.append(request)
        self.violations += audit(request, self.golden.pii_values)
        document = json.loads(request.evidence_json)
        expected = document["expected_domain"]
        offered = _offered(document, expected)
        truth = truth_fields(self.golden, offered)

        def field_of(header: str) -> str | None:
            for cell, name in truth.items():
                if header_matches((cell,), header):
                    return name
            return None

        used: set[str] = set()
        columns = []
        for column in document["columns"]:
            if column["assigned"] is not None or column["pii"] is not None or column["values"] <= 0:
                continue
            name = field_of(str(column["header"]))
            if name is not None and name not in used:
                used.add(name)
                columns.append({"index": column["index"], "field": name, "confidence": "MEDIUM"})
        answers = []
        for question in document["questions"]:
            choice: str | None = None
            if question["kind"] == "DOMAIN":
                domain = self.golden.domain.upper()
                choice = domain if domain in question["choices"] else None
            elif question["kind"] == "COLUMN_MEANING" and question["column"] is not None:
                header = next(
                    (c["header"] for c in document["columns"] if c["index"] == question["column"]),
                    "",
                )
                choice = field_of(str(header))
            elif question["kind"] == "GRID_HOURS":
                choice = self.golden.answers.get("grid_hours")
            if choice in question["choices"]:
                answers.append({"question": question["id"], "choice": choice})
        domain = self.golden.domain.upper()
        self.answered.append(len(columns) + len(answers))
        return {
            "domain": domain if domain in document["allowed"]["domains"] else "UNKNOWN",
            "columns": columns,
            "answers": answers,
        }


_DATE_WORDS = ("date", "check_in", "check_out", "booked_at", "due")
_NUMBER_WORDS = (
    "hours",
    "revenue",
    "amount",
    "total",
    "rooms",
    "guests",
    "nights",
    "quantity",
    "rate",
    "price",
    "cost",
    "fee",
    "commission",
    "net",
    "gross",
    "vat",
)


def _kind_of_field(name: str) -> str:
    if any(word in name for word in _DATE_WORDS):
        return "DATE"
    if any(word in name for word in _NUMBER_WORDS):
        return "NUMBER"
    return "TEXT"


@dataclass
class PlausibleWrong:
    """Wrong in the way most likely to get past the value checks: for every column it is asked
    about it names a field of the same KIND (a date for a date, an amount for an amount), never the
    right one, preferring the fields a business decision rests on. Claims MEDIUM."""

    golden: Golden
    requests: list[ImportUnderstandingRequest] = field(default_factory=list)
    violations: list[str] = field(default_factory=list)

    def suggest(self, request: ImportUnderstandingRequest) -> object:
        self.requests.append(request)
        self.violations += audit(request, self.golden.pii_values)
        document = json.loads(request.evidence_json)
        expected = document["expected_domain"]
        offered = _offered(document, expected)
        truth = truth_fields(self.golden, offered)
        critical = set(CRITICAL_FIELDS.get(expected.lower(), ()))
        taken = {c["assigned"] for c in document["columns"] if c["assigned"]}
        columns = []
        used: set[str] = set(taken)
        for column in document["columns"]:
            if column["assigned"] is not None or column["pii"] is not None or column["values"] <= 0:
                continue
            right = next(
                (n for cell, n in truth.items() if header_matches((cell,), str(column["header"]))),
                None,
            )
            pool = [
                f
                for f in offered
                if f not in used and f != right and _kind_of_field(f) == column["kind"]
            ]
            pool.sort(key=lambda f: (f not in critical and f.split("_")[0] not in critical, f))
            if pool:
                used.add(pool[0])
                columns.append({"index": column["index"], "field": pool[0], "confidence": "MEDIUM"})
        return {"domain": expected, "columns": columns, "answers": []}


# --- a second population: the model is actually needed ---


def variants(docs: list[Doc]) -> list[Doc]:
    """Documents the rules read completely, with ONE column the rules cannot read any more: the
    header of a field a business decision rests on is replaced by an opaque label. The truth is the
    same, except that the file is no longer allowed to be silently automatic."""
    made: list[Doc] = []
    for doc in docs:
        if doc.golden.outcome != "READY" or "xlsx" not in doc.tags:
            continue
        domain = {"labor": "labor", "bookings": "bookings", "costs": "costs"}[doc.card]
        for name in CRITICAL_FIELDS[domain]:
            source = doc.golden.fields.get(name)
            if source is None or source.kind != "COLUMN" or len(source.header) != 1:
                continue
            if source.occurrence != 1:
                continue
            content = _rename_header(doc, source.header[0])
            if content is None:
                continue
            golden = golden_from_json(manifest([doc])[0]["golden"])
            golden.fields[name] = Source("COLUMN", (UNREADABLE,))
            golden.outcome = "NOT_AUTOMATIC"
            made.append(
                Doc(
                    id=f"{doc.id}~hidden-{name}",
                    card=doc.card,
                    filename=doc.filename,
                    content=content,
                    golden=golden,
                    tags=(*doc.tags, "ai-variant"),
                    structure=doc.structure,
                    sheet_count=doc.sheet_count,
                )
            )
            break
    return made


def _rename_header(doc: Doc, header: str) -> bytes | None:
    book = load_workbook(BytesIO(doc.content))
    names = [n for n in doc.golden.sheets if n in book.sheetnames] or [book.sheetnames[0]]
    sheet = book[names[0]]
    wanted = normalize(header).strip().casefold()
    hits = []
    for row_number in doc.golden.header_rows or (1,):
        for cell in sheet[row_number]:
            if isinstance(cell.value, str) and normalize(cell.value).strip().casefold() == wanted:
                hits.append(cell)
    if len(hits) != 1:
        return None
    hits[0].value = UNREADABLE
    out = BytesIO()
    book.save(out)
    return out.getvalue()


# --- running a population with a stand-in ---


class RecordingResolver(ImportUnderstandingResolver):
    """The real resolver, remembering each resolution it makes (counts and a reason, no content)."""

    def __init__(self, provider: Any) -> None:
        super().__init__(provider)
        self.resolutions: list[AiResolution] = []

    def resolve(
        self, evidence: UnderstandingEvidence, fit_of: FitOf, guard: EgressGuard | None = None
    ) -> AiResolution:
        resolution = super().resolve(evidence, fit_of, guard)
        self.resolutions.append(resolution)
        return resolution


class Tap:
    """Wraps a stand-in and keeps what it was asked and what it answered."""

    def __init__(self, inner: Any) -> None:
        self.inner = inner
        self.requests: list[ImportUnderstandingRequest] = []
        self.responses: list[object] = []

    def suggest(self, request: ImportUnderstandingRequest) -> object:
        self.requests.append(request)
        try:
            response = self.inner.suggest(request)
        except BaseException as error:
            self.responses.append(("RAISED", type(error).__name__))
            raise
        self.responses.append(response)
        return response

    @property
    def violations(self) -> list[str]:
        return list(getattr(self.inner, "violations", []))

    def proposed_columns(self) -> int:
        """How many columns the model proposed a field for, in the answers it gave."""
        return sum(
            len(response["columns"])
            for response in self.responses
            if isinstance(response, dict) and isinstance(response.get("columns"), list)
        )


@dataclass
class Result:
    observed: Observed
    provider: Tap | None
    resolver: RecordingResolver | None = None


def _service_with(provider: Any) -> tuple[UnderstandingService, RecordingResolver]:
    tenant = TenantContext(UUID(int=1))
    session: Any = None
    resolver = RecordingResolver(provider)
    service = UnderstandingService(session, tenant, SmartImportService(session, tenant), resolver)
    return service, resolver


def _service(provider: Any) -> UnderstandingService:
    return _service_with(provider)[0]


def run_population(docs: list[Doc], make: Callable[[Doc], Any] | None) -> dict[str, Result]:
    results: dict[str, Result] = {}
    for doc in docs:
        provider = Tap(make(doc)) if make is not None else None
        service, resolver = _service_with(provider)
        observed = run(service, doc.card, doc.filename, doc.content)
        results[doc.id] = Result(observed, provider, resolver)
    return results


MODES: dict[str, Callable[[Doc], Any]] = {
    "oracle": lambda d: Oracle(d.golden),
    "plausible_wrong": lambda d: PlausibleWrong(d.golden),
    "wrong": lambda d: ScenarioProvider(Scenario.WRONG, _truth(d), d.golden.pii_values),
    "sure_and_wrong": lambda d: ScenarioProvider(
        Scenario.SURE_AND_WRONG, _truth(d), d.golden.pii_values
    ),
    "obeys_injection": lambda d: ScenarioProvider(
        Scenario.OBEYS_INJECTION, _truth(d), d.golden.pii_values
    ),
    "says_another_kind": lambda d: ScenarioProvider(
        Scenario.DISAGREE, _truth(d), d.golden.pii_values
    ),
    "agrees": lambda d: ScenarioProvider(Scenario.AGREE, _truth(d), d.golden.pii_values),
    "malformed": lambda d: ScenarioProvider(Scenario.MALFORMED, _truth(d), d.golden.pii_values),
    "not_json": lambda d: ScenarioProvider(Scenario.NOT_JSON, _truth(d), d.golden.pii_values),
    "nothing": lambda d: ScenarioProvider(Scenario.EMPTY, _truth(d), d.golden.pii_values),
    "unsupported_field": lambda d: ScenarioProvider(
        Scenario.UNSUPPORTED_FIELD, _truth(d), d.golden.pii_values
    ),
    "heavy_answer": lambda d: ScenarioProvider(Scenario.HEAVY, _truth(d), d.golden.pii_values),
    "timeout": lambda d: ScenarioProvider(Scenario.TIMEOUT, _truth(d), d.golden.pii_values),
    "http_429": lambda d: ScenarioProvider(Scenario.RATE_LIMITED, _truth(d), d.golden.pii_values),
    "http_500": lambda d: ScenarioProvider(Scenario.SERVER_ERROR, _truth(d), d.golden.pii_values),
    "raw_exception": lambda d: ScenarioProvider(
        Scenario.RAW_EXCEPTION, _truth(d), d.golden.pii_values
    ),
}
# modes whose answer must leave the file EXACTLY as it is with no model
SILENT = (
    "agrees",
    "malformed",
    "not_json",
    "nothing",
    "unsupported_field",
    "heavy_answer",
    "timeout",
    "http_429",
    "http_500",
    "raw_exception",
    "sure_and_wrong",
    "obeys_injection",
    "says_another_kind",
)


def _truth(doc: Doc) -> dict[str, str]:
    fields = {f for sem in SEMANTICS.values() for f in sem.fields}
    return truth_fields(doc.golden, sorted(fields))


# --- measuring ---


def _pct(values: list[float], q: float) -> float:
    ordered = sorted(values)
    return ordered[min(len(ordered) - 1, max(0, round(q * (len(ordered) - 1))))] if ordered else 0.0


def _stats(values: list[float]) -> dict[str, float]:
    if not values:
        return {"mean": 0.0, "p50": 0.0, "p95": 0.0, "max": 0.0}
    return {
        "mean": round(statistics.fmean(values), 1),
        "p50": statistics.median(values),
        "p95": _pct(values, 0.95),
        "max": max(values),
    }


def score_all(docs: list[Doc], results: dict[str, Result]) -> dict[str, CaseResult]:
    entries = {e["id"]: e for e in json.loads(json.dumps(manifest(docs)))}
    return {
        doc.id: score_case(entries[doc.id], golden_from_json(entries[doc.id]["golden"]), r.observed)
        for doc in docs
        for r in [results[doc.id]]
    }


def _resolutions_of(result: Result) -> list[AiResolution]:
    return result.resolver.resolutions if result.resolver is not None else []


def _proposed_of(result: Result) -> int:
    return result.provider.proposed_columns() if result.provider is not None else 0


def _requests_of(result: Result) -> list[ImportUnderstandingRequest]:
    return result.provider.requests if result.provider is not None else []


def _severe(case: CaseResult) -> set[tuple[str, str, str]]:
    return {(f.severity, f.category, f.detail) for f in case.findings if f.severity in ("P0", "P1")}


def _question_keys(observed: Observed) -> dict[tuple[str, str | None], str | None]:
    return {(q.kind, q.column): q.candidate for q in observed.questions}


def measure(
    docs: list[Doc], results: dict[str, Result], base: dict[str, Result], mode: str
) -> dict[str, Any]:
    scored = score_all(docs, results)
    base_scored = score_all(docs, base)
    n = len(docs)
    calls_per_file = [len(_requests_of(results[d.id])) for d in docs]
    asked_calls = [c for c in calls_per_file if c]
    resolutions = [(d, r) for d in docs for r in _resolutions_of(results[d.id])]
    reasons = Counter(r.reason.value for _, r in resolutions)
    payload_chars = [float(r.payload_chars) for _, r in resolutions if r.payload_chars]
    input_tokens = [float(r.input_tokens) for _, r in resolutions if r.payload_chars]
    response_chars = [float(r.response_chars) for _, r in resolutions if r.response_chars]
    blocked = [d for d, r in resolutions if r.reason == ResolutionReason.BLOCKED]
    violations = sum(len(getattr(results[d.id].provider, "violations", [])) for d in docs)
    leaks = sum(1 for d in docs if getattr(results[d.id].provider, "violations", []))
    changed = [d.id for d in docs if results[d.id].observed.payload != base[d.id].observed.payload]
    newly_automatic = [
        d.id for d in docs if results[d.id].observed.automatic and not base[d.id].observed.automatic
    ]
    raised_to_ready = [
        d.id
        for d in docs
        if results[d.id].observed.level == "READY" and base[d.id].observed.level != "READY"
    ]
    transitions = Counter(
        f"{base[d.id].observed.level}->{results[d.id].observed.level}"
        for d in docs
        if base[d.id].observed.level != results[d.id].observed.level
    )
    offered = [f for sem in SEMANTICS.values() for f in sem.fields]
    suggestions = wrong = right = critical_wrong = 0
    wrong_ids: list[str] = []
    for doc in docs:
        truth = truth_fields(doc.golden, offered)
        for claim in results[doc.id].observed.claims:
            if claim.kind != "ASSISTANT":
                continue
            suggestions += 1
            expected = next(
                (n_ for cell, n_ in truth.items() if header_matches((cell,), claim.header or "")),
                None,
            )
            good = expected is not None and (
                expected == claim.field or claim.field in PSEUDO.get(_pseudo_of(expected), ())
            )
            if good:
                right += 1
            else:
                wrong += 1
                wrong_ids.append(doc.id)
                critical = CRITICAL_FIELDS.get(doc.golden.domain, ())
                if claim.field in critical or any(
                    claim.field in PSEUDO.get(c, ()) for c in critical
                ):
                    critical_wrong += 1
    proposed = sum(_proposed_of(results[d.id]) for d in docs)
    precompiled = 0
    for d in docs:
        before, after = _question_keys(base[d.id].observed), _question_keys(results[d.id].observed)
        precompiled += sum(
            1
            for key, candidate in after.items()
            if candidate is not None and before.get(key) is None
        )
    findings = [f for c in scored.values() for f in c.findings]
    ai_fa = [
        d.id
        for d in docs
        if scored[d.id].false_automation and not base_scored[d.id].false_automation
    ]
    ai_p0 = sum(
        1
        for d in docs
        for sev, *_ in _severe(scored[d.id]) - _severe(base_scored[d.id])
        if sev == "P0"
    )
    ai_p1 = sum(
        1
        for d in docs
        for sev, *_ in _severe(scored[d.id]) - _severe(base_scored[d.id])
        if sev == "P1"
    )
    levels = Counter(r.observed.level for r in results.values())
    return {
        "mode": mode,
        "documents": n,
        "documents_where_the_model_was_asked": len(asked_calls),
        "calls": {
            "total": sum(calls_per_file),
            "max_per_file": max(calls_per_file, default=0),
            "mean_per_file": round(sum(calls_per_file) / n, 3) if n else 0,
            "p95_per_file": _pct([float(c) for c in calls_per_file], 0.95),
            "mean_per_asked_file": round(sum(asked_calls) / len(asked_calls), 3)
            if asked_calls
            else 0,
            "p95_per_asked_file": _pct([float(c) for c in asked_calls], 0.95),
        },
        "resolutions_by_reason": dict(reasons),
        "egress": {
            "payload_chars": _stats(payload_chars),
            "approx_input_tokens": _stats(input_tokens),
            "response_chars": _stats(response_chars),
            "hard_limit_chars": MAX_PAYLOAD_CHARS,
            "over_the_limit": sum(1 for v in payload_chars if v > MAX_PAYLOAD_CHARS),
            "reduced_to_fit": sum(1 for _, r in resolutions if r.reduced),
            "blocked_calls": len(blocked),
            "blocked_documents": sorted({d.id for d in blocked}),
            "false_positive_blocked_calls": sum(1 for d in blocked if not d.golden.pii_values),
            "actual_pii_leaks": violations,
            "documents_with_an_actual_pii_leak": leaks,
        },
        "payload_chars": _stats(payload_chars),
        "approx_input_tokens": _stats(input_tokens),
        "levels": dict(levels),
        "level_transitions": dict(transitions),
        "automatic": sum(1 for r in results.values() if r.observed.automatic),
        "automatic_at_baseline": sum(1 for r in base.values() if r.observed.automatic),
        "documents_whose_screen_changed": len(changed),
        "newly_automatic": newly_automatic,
        "raised_to_ready": raised_to_ready,
        "false_automation": sum(1 for c in scored.values() if c.false_automation),
        "false_automation_at_baseline": sum(1 for c in base_scored.values() if c.false_automation),
        "ai_caused": {
            "false_automation_documents": ai_fa,
            "new_p0": ai_p0,
            "new_p1": ai_p1,
            "automatic_imports": len(newly_automatic),
        },
        "p0": sum(1 for f in findings if f.severity == "P0"),
        "p1": sum(1 for f in findings if f.severity == "P1"),
        "browser_pii_leaked_values": sum(c.pii_values_leaked for c in scored.values()),
        "model_payload_violations": violations,
        "documents_with_model_payload_violations": leaks,
        "safe_fallback": ratio(
            sum(1 for c in scored.values() if c.safe_fallback is True),
            sum(1 for c in scored.values() if c.safe_fallback is not None),
        ),
        "necessary_question_recall": ratio(
            sum(c.questions_necessary_found for c in scored.values()),
            sum(len(c.questions_expected) for c in scored.values()),
        ),
        "questions_asked": sum(len(c.questions_asked) for c in scored.values()),
        "questions_asked_at_baseline": sum(len(c.questions_asked) for c in base_scored.values()),
        "questions_unnecessary": sum(c.questions_unnecessary for c in scored.values()),
        "model_suggestions": {
            "proposed_columns": proposed,
            "accepted_as_medium": suggestions,
            "discarded": max(0, proposed - suggestions),
            "right": right,
            "wrong": wrong,
            "wrong_on_critical_fields": critical_wrong,
            "documents_with_a_wrong_suggestion": len(set(wrong_ids)),
            "questions_precompiled": precompiled,
            "all_at_medium": all(
                c.level == "MEDIUM"
                for r in results.values()
                for c in r.observed.claims
                if c.kind == "ASSISTANT"
            ),
        },
        "ids": {"changed": changed},
    }


def _pseudo_of(name: str) -> str:
    for pseudo, aliases in PSEUDO.items():
        if name in aliases:
            return pseudo
    return name


def value_of_a_perfect_model(
    docs: list[Doc], oracle: dict[str, Result], base: dict[str, Result]
) -> dict[str, Any]:
    """What a perfect answer would be worth, document by document."""
    rows: list[dict[str, Any]] = []
    for doc in docs:
        a, b = base[doc.id].observed, oracle[doc.id].observed
        if not _requests_of(oracle[doc.id]):
            continue
        before, after = _question_keys(a), _question_keys(b)
        rows.append(
            {
                "id": doc.id,
                "truth": doc.golden.outcome,
                "baseline_level": a.level,
                "with_a_perfect_model_level": b.level,
                "baseline_questions": a.question_kinds,
                "with_a_perfect_model_questions": b.question_kinds,
                "improved": b.payload != a.payload,
                "suggestions": sum(1 for c in b.claims if c.kind == "ASSISTANT"),
                "questions_precompiled": sum(
                    1 for k, cand in after.items() if cand is not None and before.get(k) is None
                ),
                "automatic_with_the_model": b.automatic,
                "automatic_at_baseline": a.automatic,
            }
        )
    return {
        "documents_the_model_was_asked_about": len(rows),
        "documents_really_improved": sum(1 for r in rows if r["improved"]),
        "questions_precompiled": sum(r["questions_precompiled"] for r in rows),
        "suggestions_accepted_at_medium": sum(r["suggestions"] for r in rows),
        "manual_to_ask": sum(
            1
            for r in rows
            if r["baseline_level"] == "MANUAL" and r["with_a_perfect_model_level"] == "ASK"
        ),
        "blocked_to_ask": sum(
            1
            for r in rows
            if r["baseline_level"] == "BLOCKED" and r["with_a_perfect_model_level"] == "ASK"
        ),
        "ask_to_ask_with_a_confirmation_instead_of_an_open_question": sum(
            1
            for r in rows
            if r["baseline_level"] == "ASK"
            and r["with_a_perfect_model_level"] == "ASK"
            and r["improved"]
        ),
        "new_automatic_imports": sum(
            1 for r in rows if r["automatic_with_the_model"] and not r["automatic_at_baseline"]
        ),
        "documents": rows,
    }


# --- the 21 questions that remain (see the report for the reasoning) ---

# class: AI_HELPFUL | DETERMINISTIC_REQUIRED | HUMAN_REQUIRED | UNSUPPORTED
RESIDUAL = [
    (
        "L36_abbreviations_known",
        "VALUE_LABELS F&B, FO, HK, HSKP, REC",
        "UNSUPPORTED",
        "Hotel jargon a model reads - but the closed answer has no 'meaning of a label'; a lexicon "
        "would also do it. Codes stay asked by product decision.",
    ),
    (
        "L62_neg_swapped_header_labels",
        "DEPARTMENT_ROWS",
        "DETERMINISTIC_REQUIRED",
        "The headers contradict the values: a contradiction is the rules' to detect, never a "
        "model's to smooth over.",
    ),
    (
        "LO07_department_in_a_label_value_block",
        "DEPARTMENT_ROWS",
        "DETERMINISTIC_REQUIRED",
        "'Reparto: X' above the table is structure the parser should read; no model needed.",
    ),
    (
        "LO15_abbreviated_department_header_and_codes",
        "VALUE_LABELS FO, HK, REC",
        "UNSUPPORTED",
        "As L36: a label-meaning capability, or a lexicon. Codes stay asked by product decision.",
    ),
    (
        "B14_booker_and_guest_names_two_columns",
        "COLUMN_MEANING Price",
        "HUMAN_REQUIRED",
        "Which amount 'Price' is, is not in the file; money is never guessed (product decision).",
    ),
    (
        "B41_office_channel_manager_extract",
        "COLUMN_MEANING Total",
        "HUMAN_REQUIRED",
        "A bare 'Total': room revenue or the whole stay? Approved to stay MEDIUM + confirm.",
    ),
    (
        "B42_office_front_desk_monthly_report",
        "COLUMN_MEANING Ricavo",
        "HUMAN_REQUIRED",
        "Which revenue is not in the file; money is never guessed.",
    ),
    (
        "B50_neg_arrival_and_departure_swapped",
        "DATES_REVERSED Arrivo",
        "DETERMINISTIC_REQUIRED",
        "Every arrival after its departure: a value check, and a person decides.",
    ),
    (
        "B51_neg_occupancy_calendar_not_bookings",
        "DOMAIN",
        "DETERMINISTIC_REQUIRED",
        "A room-by-day calendar is not bookings: the kind of data is the rules' call.",
    ),
    (
        "B55_neg_checkout_before_arrival_rows",
        "DATES_REVERSED Arrivo",
        "DETERMINISTIC_REQUIRED",
        "Departures before arrivals: a value check.",
    ),
    (
        "B64_adv_near_identical_date_headers",
        "COLUMN_MEANING Check-in",
        "HUMAN_REQUIRED",
        "Planned and actual check-in side by side: two candidates for one field.",
    ),
    (
        "B64_adv_near_identical_date_headers",
        "COLUMN_MEANING Check-in effettivo",
        "HUMAN_REQUIRED",
        "Same file: the choice between the two is the hotel's.",
    ),
    (
        "BO01_channel_written_as_codes",
        "COLUMN_MEANING Imp.",
        "HUMAN_REQUIRED",
        "An abbreviated amount: which one is not in the file.",
    ),
    (
        "BO03_channel_manager_abbreviations",
        "COLUMN_MEANING Tot",
        "HUMAN_REQUIRED",
        "An abbreviated total: which one is not in the file.",
    ),
    (
        "BO09_status_written_as_ok_canc_ns",
        "VALUE_LABELS OK, CANC, NS",
        "UNSUPPORTED",
        "Status codes: a label-meaning capability or a lexicon.",
    ),
    (
        "BO09_status_written_as_ok_canc_ns",
        "COLUMN_MEANING Stato",
        "AI_HELPFUL",
        "The rules already propose 'status' (MEDIUM): a model would only agree, and agreeing "
        "cannot lift MEDIUM.",
    ),
    (
        "C33_credit_notes_negative_amounts",
        "COLUMN_MEANING Imponibile",
        "HUMAN_REQUIRED",
        "Negative amounts: credit note or error - the hotel's convention.",
    ),
    (
        "C38_several_lines_per_invoice",
        "COLUMN_MEANING N. documento",
        "DETERMINISTIC_REQUIRED",
        "The same number on several lines is how invoices are written: a rule, not a model.",
    ),
    (
        "C44_office_accounting_export_header_on_row_five",
        "COLUMN_MEANING Data reg.",
        "HUMAN_REQUIRED",
        "Registration date or invoice date: an accounting convention not in the file.",
    ),
    (
        "C46_office_manual_expense_sheet",
        "COLUMN_MEANING Da chi",
        "AI_HELPFUL",
        "Informal wording ('from whom'): exactly what a model reads - but the rules already "
        "propose it at MEDIUM.",
    ),
    (
        "C46_office_manual_expense_sheet",
        "COLUMN_MEANING Cosa",
        "AI_HELPFUL",
        "Informal wording ('what'): as above.",
    ),
]


def classify_residual(results_oracle: dict[str, Result]) -> dict[str, Any]:
    counts = Counter(row[2] for row in RESIDUAL)
    rows = []
    for doc_id, question, klass, why in RESIDUAL:
        result = results_oracle.get(doc_id)
        calls = len(getattr(result.provider, "requests", [])) if result else None
        rows.append(
            {
                "document": doc_id,
                "question": question,
                "class": klass,
                "why": why,
                "the_model_is_even_asked_about_this_file": bool(calls),
            }
        )
    return {
        "total": len(RESIDUAL),
        "by_class": dict(counts),
        "ai_could_remove_today": 0,
        "rows": rows,
    }


# --- the report ---

SCOPE = (
    "FIRST SCREEN ONLY: each document is understood once, as a person's first upload, and scored "
    "against the truth WITHOUT the conversation rounds. P0/P1 and the other counts here are NOT "
    "comparable with the official full-flow benchmark (docs/benchmarks/"
    "smart-import-v3-question-reduction-1.md), which also plays the answers a careful person "
    "gives. "
    "What this benchmark compares is each mode against the SAME run with no model."
)


def markdown(data: dict[str, Any], today: str) -> str:
    """The written report. `modes[...]` holds one entry per stand-in AND a small `no_model` entry
    (the reference run); only the stand-ins have the full set of measures."""
    lines = [
        "# Smart Import V3 - AI readiness benchmark",
        "",
        f"Run {today} on branch `feature/smart-import-v3-ai-readiness` "
        f"(base `main` {BASELINE[:7]}). "
        "The model is a stand-in (an oracle, models that are wrong, models that fail); nothing "
        "leaves the process and no vendor is involved.",
        "",
        f"> **Scope.** {SCOPE}",
        "",
        "## What was run",
        "",
        f"* **corpus** - the {data['corpus']['documents']} documents of the Smart Import benchmark "
        "(hash unchanged).",
        f"* **variants** - {data['variants']['documents']} documents the rules read completely, "
        "each with one decisive column whose header was replaced by an opaque label, so that the "
        "model is actually asked. Same files, same truth, except that the file may not be "
        "automatic any more.",
        "",
        "## Safety with a model that is wrong, sure of it, obedient or broken",
        "",
        "AI-caused = compared with the run with no model. "
        "Newly automatic and made READY must be 0.",
        "",
        "| model | population | asked | screens changed | newly automatic | made READY "
        "| AI-caused false automation | AI-caused P0 | AI-caused P1 | PII browser | PII to model |",
        "|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for population in ("corpus", "variants"):
        for mode, row in data["modes"][population].items():
            if mode in ("oracle", "no_model"):
                continue
            caused = row["ai_caused"]
            lines.append(
                f"| {mode} | {population} | {row['documents_where_the_model_was_asked']} "
                f"| {row['documents_whose_screen_changed']} | {len(row['newly_automatic'])} "
                f"| {len(row['raised_to_ready'])} | {len(caused['false_automation_documents'])} "
                f"| {caused['new_p0']} | {caused['new_p1']} "
                f"| {row['browser_pii_leaked_values']} | {row['model_payload_violations']} |"
            )
    lines += ["", "## A wrong model: what it changes and what protects the file", ""]
    lines += [
        "| model | population | proposed | accepted (MEDIUM) | discarded | MANUAL to ASK "
        "| BLOCKED to ASK | questions added or precompiled |",
        "|---|---|---:|---:|---:|---:|---:|---:|",
    ]
    for population in ("corpus", "variants"):
        for mode in ("plausible_wrong", "wrong", "sure_and_wrong", "obeys_injection"):
            row = data["modes"][population].get(mode)
            if row is None:
                continue
            m, t = row["model_suggestions"], row["level_transitions"]
            lines.append(
                f"| {mode} | {population} | {m['proposed_columns']} | {m['accepted_as_medium']} "
                f"| {m['discarded']} | {t.get('MANUAL->ASK', 0)} | {t.get('BLOCKED->ASK', 0)} "
                f"| {row['questions_asked'] - row['questions_asked_at_baseline']:+d} |"
            )
    lines += ["", "## Cost and egress (the model is asked about the files with something open)", ""]
    lines += [
        "| population | asked | calls/file mean | p95 | max | payload mean | p50 | p95 | max "
        "| hard limit | blocked | false-positive blocks | actual PII leaks |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for population in ("corpus", "variants"):
        row = data["modes"][population]["oracle"]
        calls, egress = row["calls"], row["egress"]
        payload = egress["payload_chars"]
        lines.append(
            f"| {population} | {row['documents_where_the_model_was_asked']} of {row['documents']} "
            f"| {calls['mean_per_file']} | {calls['p95_per_file']:.0f} | {calls['max_per_file']} "
            f"| {payload['mean']:.0f} | {payload['p50']:.0f} | {payload['p95']:.0f} "
            f"| {payload['max']:.0f} | {egress['hard_limit_chars']} | {egress['blocked_calls']} "
            f"| {egress['false_positive_blocked_calls']} | {egress['actual_pii_leaks']} |"
        )
    lines += ["", "## What a perfect model would be worth", ""]
    lines += [
        "| population | asked | really improved | questions precompiled | suggestions at MEDIUM "
        "| MANUAL to ASK | BLOCKED to ASK | ASK with a confirmation | new automatic |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for population in ("corpus", "variants"):
        v = data["value"][population]
        lines.append(
            f"| {population} | {v['documents_the_model_was_asked_about']} "
            f"| {v['documents_really_improved']} | {v['questions_precompiled']} "
            f"| {v['suggestions_accepted_at_medium']} | {v['manual_to_ask']} "
            f"| {v['blocked_to_ask']} "
            f"| {v['ask_to_ask_with_a_confirmation_instead_of_an_open_question']} "
            f"| {v['new_automatic_imports']} |"
        )
    residual = data["residual"]
    lines += [
        "",
        f"## The {residual['total']} questions that remain on the corpus",
        "",
        f"By class: {residual['by_class']}. A perfect model could remove "
        f"**{residual['ai_could_remove_today']}** of them today.",
        "",
        "| document | question | class | why |",
        "|---|---|---|---|",
    ]
    for row in residual["rows"]:
        lines.append(f"| {row['document']} | {row['question']} | {row['class']} | {row['why']} |")
    return "\n".join(lines) + "\n"


def say(message: str) -> None:
    print(message, file=sys.stderr, flush=True)


def run_benchmark(corpus_docs: list[Doc], progress: Callable[[str], None] = say) -> dict[str, Any]:
    """Every mode on both populations; returns the results (nothing is written)."""
    populations = {"corpus": corpus_docs, "variants": variants(corpus_docs)}
    data: dict[str, Any] = {
        "benchmark": "smart-import-v3-ai-readiness-v1",
        "baseline_main": BASELINE,
        "assistant": "stand-in providers only; IMPORT_UNDERSTANDING_PROVIDER stays unconfigured",
        "scope": SCOPE,
        "corpus": {"documents": len(corpus_docs)},
        "variants": {"documents": len(populations["variants"])},
        "modes": {},
        "value": {},
    }
    oracles: dict[str, dict[str, Result]] = {}
    for population, docs in populations.items():
        progress(f"{population}: {len(docs)} documents - baseline (no model)")
        base = run_population(docs, None)
        data["modes"][population] = {}
        for mode, make in MODES.items():
            progress(f"{population}: {mode}")
            results = run_population(docs, make)
            data["modes"][population][mode] = measure(docs, results, base, mode)
            if mode == "oracle":
                oracles[population] = results
        data["value"][population] = value_of_a_perfect_model(docs, oracles[population], base)
        data["modes"][population]["no_model"] = {
            "questions_asked": sum(len(r.observed.question_kinds) for r in base.values()),
            "automatic": sum(1 for r in base.values() if r.observed.automatic),
        }
        if population == "corpus":
            data["residual"] = classify_residual(oracles["corpus"])
    data["summary"] = _summary(data)
    return data


def write_outputs(data: dict[str, Any], json_path: Path, markdown_path: Path, today: str) -> None:
    """JSON and Markdown, UTF-8 with LF line ends on every platform."""
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
    )
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(markdown(data, today), encoding="utf-8", newline="\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m tests.benchmarks.smart_import_v3.ai_readiness")
    parser.add_argument("--json", type=Path, default=DEFAULT_JSON)
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MD)
    parser.add_argument("--only", action="append", default=[], help="document ids containing this")
    parser.add_argument("--no-write", action="store_true")
    args = parser.parse_args(argv)

    corpus_docs = build_corpus()
    if args.only:
        corpus_docs = [d for d in corpus_docs if any(o in d.id for o in args.only)]
    data = run_benchmark(corpus_docs)
    print(json.dumps(data["summary"], indent=1))
    # a subset must never overwrite the official result: it writes only where it is told to
    official_paths = args.json == DEFAULT_JSON and args.markdown == DEFAULT_MD
    if args.no_write or (args.only and official_paths):
        return 0
    write_outputs(data, args.json, args.markdown, date.today().isoformat())
    say(f"JSON {args.json}\nMarkdown {args.markdown}")
    return 0


def _summary(data: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {"scope": "first screen only; see `scope`"}
    for population, modes in data["modes"].items():
        stand_ins = {k: m for k, m in modes.items() if k not in ("oracle", "no_model")}
        out[population] = {
            "ai_caused_false_automation": sum(
                len(m["ai_caused"]["false_automation_documents"]) for m in stand_ins.values()
            ),
            "ai_caused_new_p0": sum(m["ai_caused"]["new_p0"] for m in stand_ins.values()),
            "ai_caused_new_p1": sum(m["ai_caused"]["new_p1"] for m in stand_ins.values()),
            "newly_automatic": sum(len(m["newly_automatic"]) for m in stand_ins.values()),
            "raised_to_ready_by_a_wrong_model": sum(
                len(m["raised_to_ready"]) for m in stand_ins.values()
            ),
            "oracle_raised_to_ready": len(modes["oracle"]["raised_to_ready"]),
            "oracle_newly_automatic": len(modes["oracle"]["newly_automatic"]),
            "model_payload_violations": sum(
                m["model_payload_violations"] for m in stand_ins.values()
            ),
            "blocked_calls": sum(
                m["egress"]["blocked_calls"] for m in modes.values() if "egress" in m
            ),
            "silent_modes_that_changed_a_screen": {
                k: modes[k]["documents_whose_screen_changed"]
                for k in SILENT
                if modes[k]["documents_whose_screen_changed"]
            },
        }
    return out


if __name__ == "__main__":
    raise SystemExit(main())
